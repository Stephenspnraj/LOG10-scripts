import os
import csv
import pymysql
from datetime import datetime
from tabulate import tabulate
from colorama import init, Fore
from openpyxl import Workbook
from contextlib import contextmanager

# ------------------------------
# CONFIG
# ------------------------------
INPUT_CSV_PATH = "pending_manifest_corrections.csv"
OUTPUT_XLSX_PATH = "pending_manifest_corrections_output.xlsx"

DB_CONFIG = {
    "host": os.getenv("DB_HOST", "log10-replica-replica.cco3osxqlq4g.ap-south-1.rds.amazonaws.com"),
    "user": os.getenv("DB_USER", "log10_scripts"),
    "password": os.getenv("DB_PASSWORD", "D2lx7Wz0Wm9ISAY-Vp1gxKDTzLRRl5k1m"),
    "database": os.getenv("DB_NAME", "loadshare"),
    "port": int(os.getenv("DB_PORT", "3306")),
}

# ------------------------------
# DB Connection with Context Manager
# ------------------------------
@contextmanager
def get_connection():
    conn = pymysql.connect(
        host=DB_CONFIG["host"],
        user=DB_CONFIG["user"],
        password=DB_CONFIG["password"],
        db=DB_CONFIG["database"],
        port=DB_CONFIG["port"],
        cursorclass=pymysql.cursors.DictCursor,
        charset="utf8mb4",
    )
    try:
        yield conn
    finally:
        conn.close()

# ------------------------------
# Preload Locations
# ------------------------------
def preload_locations(conn):
    with conn.cursor() as cur:
        cur.execute("""
            SELECT id, alias, client_location_name
            FROM locations
            WHERE entity_type='PARTNER' AND status=1
        """)
        rows = cur.fetchall()

    location_map = {}
    for r in rows:
        if r["alias"]:
            location_map[r["alias"].strip().lower()] = {"id": r["id"], "original": r["alias"].strip()}
        if r["client_location_name"]:
            location_map[r["client_location_name"].strip().lower()] = {
                "id": r["id"],
                "original": r["client_location_name"].strip()
            }
    return location_map

def resolve_from_cache(location_map, value):
    if not value:
        return None, None
    stripped = value.strip()
    lowered = stripped.lower()
    entry = location_map.get(lowered)
    if entry:
        return entry["id"], entry["original"]
    return None, stripped

# ------------------------------
# Helpers
# ------------------------------
def normalize_headers(headers):
    return {h.strip().lower(): h for h in headers}

def split_destinations(value: str):
    if value is None:
        return []
    return [part.strip() for part in str(value).split(',') if part.strip()]

# ------------------------------
# Bulk Fetch Pending Manifests
# ------------------------------
def fetch_all_pending_manifests(conn, pairs):
    if not pairs:
        return []

    placeholders = ", ".join(["(%s, %s)"] * len(pairs))
    flat_values = [val for tup in pairs for val in tup]

    query = f"""
        SELECT m.manifest_code,
               cl.alias AS current_loc,
               dl.alias AS destination_loc,
               nl.alias AS next_loc,
               m.updated_at ,
                      cl.id AS current_id,
               dl.id AS dest_id       FROM manifests m
        JOIN locations cl ON cl.id = m.current_loc_id
        JOIN locations dl ON dl.id = m.destination_loc_id
        JOIN locations nl ON nl.id = m.next_loc_id
        WHERE (cl.id, dl.id) IN ({placeholders})
          AND m.manifest_status = 'PENDING'
          AND m.is_active = 1
          AND m.updated_at > NOW() - INTERVAL 30 DAY
    """

    with conn.cursor() as cur:
        cur.execute(query, flat_values)
        return cur.fetchall()

# ------------------------------
# Batch Update
# ------------------------------
def batch_update_pending_manifests(conn, update_groups):
    total_updated = 0
    summary_updates = []
    with conn.cursor() as cur:
        for (curr_id, dest_id, next_id), _ in update_groups.items():
            cur.execute(
                """
                UPDATE manifests
                   SET next_loc_id = %s
                 WHERE current_loc_id = %s
                   AND destination_loc_id = %s
                   AND manifest_status = 'PENDING'
                   AND is_active = 1
                   AND updated_at > NOW() - INTERVAL 30 DAY
                """,
                (next_id, curr_id, dest_id),
            )
            updated = cur.rowcount
            total_updated += updated
            summary_updates.append((curr_id, dest_id, next_id, updated))
            conn.commit()
    return total_updated, summary_updates

# ------------------------------
# Main
# ------------------------------
def main():
    init(autoreset=True, strip=False)

    if not os.path.exists(INPUT_CSV_PATH):
        print(Fore.RED + f"❌ CSV file not found: {INPUT_CSV_PATH}")
        return

    try:
        with get_connection() as conn:
            location_map = preload_locations(conn)

            output_rows = []
            remarks_rows = []
            total_selected = 0
            total_updated = 0
            update_groups = {}
            update_map_for_remarks = {}
            next_loc_map = {}

            # First pass: collect all pairs to bulk fetch
            pairs_to_fetch = []
            csv_rows = []  # keep raw rows for second pass

            with open(INPUT_CSV_PATH, newline='') as f:
                reader = csv.DictReader(f)
                headers_map = normalize_headers(reader.fieldnames or [])
                required = ["current_location", "destination_locations", "next_location"]
                missing_cols = [c for c in required if c.lower() not in headers_map]
                if missing_cols:
                    print(Fore.RED + f"❌ Missing required columns: {', '.join(missing_cols)}")
                    return

                for row_num, row in enumerate(reader, start=1):
                    csv_rows.append((row_num, row))  # store for later processing

                    current_in = (row.get(headers_map["current_location"], "") or "").strip()
                    destin_in = (row.get(headers_map["destination_locations"], "") or "").strip()
                    next_in = (row.get(headers_map["next_location"], "") or "").strip()

                    curr_id, _ = resolve_from_cache(location_map, current_in)
                    next_id, _ = resolve_from_cache(location_map, next_in)
                    if not curr_id or not next_id:
                        continue

                    dest_list = split_destinations(destin_in)
                    for dest_in in dest_list:
                        dest_id, _ = resolve_from_cache(location_map, dest_in)
                        if dest_id:
                            pairs_to_fetch.append((curr_id, dest_id))

            # Bulk fetch manifests
            all_pending = fetch_all_pending_manifests(conn, pairs_to_fetch)

            # Organize results by (curr_id, dest_id)
            pending_map = {}
            for r in all_pending:
                key = (r["current_id"], r["dest_id"])
                pending_map.setdefault(key, []).append(r)

            # Second pass: process rows
            for row_num, row in csv_rows:
                current_in = (row.get(headers_map["current_location"], "") or "").strip()
                destin_in = (row.get(headers_map["destination_locations"], "") or "").strip()
                next_in = (row.get(headers_map["next_location"], "") or "").strip()

                print(Fore.CYAN + "\n" + "=" * 60)
                print(Fore.CYAN + f"Processing Row {row_num}")
                print(Fore.CYAN + f"Current: {current_in} | Destinations: {destin_in} | Next: {next_in}")
                print(Fore.CYAN + "=" * 60)

                curr_id, curr_display = resolve_from_cache(location_map, current_in)
                next_id, next_display = resolve_from_cache(location_map, next_in)

                if not curr_id:
                    remarks_rows.append({
                        "current_location": current_in,
                        "destination": destin_in,
                        "next_location": next_in,
                        "remarks": f"Missing: current_location='{current_in}'",
                    })
                    print(Fore.YELLOW + f"⚠ Skipping row: invalid current location '{current_in}'")
                    continue
                if not next_id:
                    remarks_rows.append({
                        "current_location": current_in,
                        "destination": destin_in,
                        "next_location": next_in,
                        "remarks": f"Missing: next_location='{next_in}'",
                    })
                    print(Fore.YELLOW + f"⚠ Skipping row: invalid next location '{next_in}'")
                    continue

                dest_list = split_destinations(destin_in)
                if not dest_list:
                    remarks_rows.append({
                        "current_location": current_in,
                        "destination": "",
                        "next_location": next_in,
                        "remarks": "No destination provided",
                    })
                    print(Fore.YELLOW + f"⚠ Skipping row: destination_locations empty")
                    continue

                for dest_in in dest_list:
                    dest_id, dest_display = resolve_from_cache(location_map, dest_in)
                    if not dest_id:
                        remarks_rows.append({
                            "current_location": current_in,
                            "destination": dest_in,
                            "next_location": next_in,
                            "remarks": f"Missing: destination_locations='{dest_in}'",
                        })
                        print(Fore.YELLOW + f"⚠ Skipping destination '{dest_in}': invalid location")
                        continue

                    key = (curr_id, dest_id)
                    pending = pending_map.get(key, [])
                    print(f"📦 Pending manifests for {curr_display} → {dest_display}: {len(pending)} found")

                    if pending:
                        print(tabulate(pending, headers="keys", tablefmt="pretty"))
                        total_selected += len(pending)
                        for r in pending:
                            output_rows.append({
                                "manifest_code": r["manifest_code"],
                                "current_loc": curr_display,
                                "destination_loc": dest_display,
                                "next_loc": r["next_loc"],
                                "updated_at": r["updated_at"].strftime('%Y-%m-%d %H:%M:%S'),
                                "expected_next_loc": next_display,
                            })

                    update_key = (curr_id, dest_id, next_id)
                    update_groups[update_key] = update_groups.get(update_key, 0) + 1
                    update_map_for_remarks[update_key] = update_map_for_remarks.get(update_key, 0) + len(pending)
                    next_loc_map[update_key] = next_display

            # Perform batch updates
            updated_total, summary_updates_raw = batch_update_pending_manifests(conn, update_groups)
            total_updated += updated_total

            # Build remarks and summary updates
            summary_updates = []
            for curr_id, dest_id, next_id, updated in summary_updates_raw:
                curr_display = next((entry["original"] for k, entry in location_map.items() if entry["id"] == curr_id), str(curr_id))
                dest_display = next((entry["original"] for k, entry in location_map.items() if entry["id"] == dest_id), str(dest_id))
                next_display = next_loc_map.get((curr_id, dest_id, next_id), "")
                count_pending = update_map_for_remarks.get((curr_id, dest_id, next_id), 0)
                remarks_rows.append({
                    "current_location": curr_display,
                    "destination": dest_display,
                    "next_location": next_display,
                    "remarks": f"Updated {updated} manifests (Pending before update: {count_pending})",
                })
                summary_updates.append({
                    "Current_location": curr_display,
                    "Destination": dest_display,
                    "No of manifests updated": updated,
                })

            # Save to Excel
            wb = Workbook()
            ws1 = wb.active
            ws1.title = "Remarks"
            ws1.append(["current_location", "destination", "next_location", "remarks"])
            for r in remarks_rows:
                ws1.append([r.get("current_location", ""), r.get("destination", ""), r.get("next_location", ""), r.get("remarks", "")])

            ws2 = wb.create_sheet(title="Pending Manifests")
            ws2.append(["manifest_code", "current_loc", "destination_loc", "next_loc", "updated_at", "expected_next_loc"])
            for r in output_rows:
                ws2.append([
                    r.get("manifest_code", ""),
                    r.get("current_loc", ""),
                    r.get("destination_loc", ""),
                    r.get("next_loc", ""),
                    r.get("updated_at", ""),
                    r.get("expected_next_loc", ""),
                ])

            wb.save(OUTPUT_XLSX_PATH)

            print(Fore.CYAN + "\nSUMMARY:")
            print(Fore.GREEN + f"✅ Total pending manifests listed: {total_selected}")
            print(Fore.GREEN + f"✅ Total manifests updated: {total_updated}")

            if summary_updates:
                print("\nPer-destination update summary:")
                print(tabulate(summary_updates, headers="keys", tablefmt="pretty"))
            else:
                print("\n(No updates performed for any (current, destination) pairs)")

            print(Fore.GREEN + f"\n✅ Results written to: {OUTPUT_XLSX_PATH}")

    except pymysql.MySQLError as e:
        print(Fore.RED + f"❌ Database error: {e}")
    except Exception as e:
        print(Fore.RED + f"❌ Unexpected error: {e}")

if __name__ == "__main__":
    main()
